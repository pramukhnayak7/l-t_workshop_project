"""Local document ingestion and lightweight semantic retrieval."""

from __future__ import annotations

import hashlib
import json
import math
import re
import sqlite3
from pathlib import Path
from typing import Iterable

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter


DATA_DIR = Path(__file__).resolve().parents[1] / "data"
DATABASE_PATH = DATA_DIR / "campus_compass.sqlite3"
EMBEDDING_SIZE = 384
STOP_WORDS = {
    "a", "about", "after", "all", "also", "an", "and", "any", "are", "as", "at",
    "be", "because", "before", "but", "by", "can", "could", "do", "does", "for",
    "from", "had", "has", "have", "how", "i", "if", "in", "into", "is", "it", "its",
    "may", "me", "my", "of", "on", "or", "our", "please", "should", "so", "student",
    "students", "that", "the", "their", "them", "there", "this", "to", "what", "when",
    "where", "which", "who", "why", "will", "with", "would", "you", "your",
}


def _tokens(text: str) -> list[str]:
    words = [
        word
        for word in re.findall(r"[a-z0-9]+", text.lower())
        if word not in STOP_WORDS
    ]
    return words + [f"{left}_{right}" for left, right in zip(words, words[1:])]


def embed_text(text: str) -> list[float]:
    """Create a deterministic, dependency-free vector for local retrieval."""
    vector = [0.0] * EMBEDDING_SIZE
    for token in _tokens(text):
        digest = hashlib.blake2b(token.encode("utf-8"), digest_size=8).digest()
        index = int.from_bytes(digest[:4], "little") % EMBEDDING_SIZE
        vector[index] += 1.0
    magnitude = math.sqrt(sum(value * value for value in vector))
    return [value / magnitude for value in vector] if magnitude else vector


def _connect() -> sqlite3.Connection:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(DATABASE_PATH)
    connection.execute(
        """CREATE TABLE IF NOT EXISTS document_chunks (
            id INTEGER PRIMARY KEY,
            source TEXT NOT NULL,
            content TEXT NOT NULL,
            embedding TEXT NOT NULL,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        )"""
    )
    return connection


def split_text(text: str, source: str) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=900,
        chunk_overlap=140,
        add_start_index=True,
    )
    return splitter.create_documents([text], metadatas=[{"source": source}])


def add_documents(documents: Iterable[Document]) -> int:
    records = [
        (doc.metadata.get("source", "Uploaded document"), doc.page_content, json.dumps(embed_text(doc.page_content)))
        for doc in documents
        if doc.page_content.strip()
    ]
    if not records:
        return 0
    with _connect() as connection:
        sources = {record[0] for record in records}
        connection.executemany(
            "DELETE FROM document_chunks WHERE source = ?",
            [(source,) for source in sources],
        )
        connection.executemany(
            "INSERT INTO document_chunks (source, content, embedding) VALUES (?, ?, ?)",
            records,
        )
    return len(records)


def ingest_text(text: str, source: str) -> int:
    return add_documents(split_text(text, source))


def ingest_file(filename: str, content: bytes) -> int:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        import io

        reader = PdfReader(io.BytesIO(content))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif suffix == ".docx":
        from docx import Document as WordDocument

        import io

        word_doc = WordDocument(io.BytesIO(content))
        text = "\n".join(paragraph.text for paragraph in word_doc.paragraphs)
    else:
        text = content.decode("utf-8", errors="replace")
    if not text.strip():
        raise ValueError(
            "No text could be extracted. Scanned PDFs need OCR before they can be indexed."
        )
    chunks = ingest_text(text, filename)
    if chunks == 0:
        raise ValueError("The document did not contain any indexable text.")
    return chunks


def retrieve(query: str, limit: int = 5) -> list[Document]:
    query_vector = embed_text(query)
    with _connect() as connection:
        rows = connection.execute(
            "SELECT source, content, embedding FROM document_chunks"
        ).fetchall()
    ranked: list[tuple[float, str, str]] = []
    for source, content, serialized_vector in rows:
        vector = json.loads(serialized_vector)
        score = sum(left * right for left, right in zip(query_vector, vector))
        if score > 0:
            ranked.append((score, source, content))
    ranked.sort(key=lambda item: item[0], reverse=True)
    return [
        Document(page_content=content, metadata={"source": source, "score": score})
        for score, source, content in ranked[:limit]
    ]


def list_sources() -> list[dict[str, str | int]]:
    with _connect() as connection:
        rows = connection.execute(
            "SELECT source, COUNT(*) FROM document_chunks GROUP BY source ORDER BY source"
        ).fetchall()
    return [{"source": source, "chunks": count} for source, count in rows]


def source_preview(source: str) -> str | None:
    with _connect() as connection:
        row = connection.execute(
            "SELECT content FROM document_chunks WHERE source = ? ORDER BY id LIMIT 1",
            (source,),
        ).fetchone()
    return row[0] if row else None


def delete_source(source: str) -> None:
    with _connect() as connection:
        connection.execute("DELETE FROM document_chunks WHERE source = ?", (source,))